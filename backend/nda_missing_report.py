"""Build the back-office gap report: new clients the CRM has that the NDA Excel sheets missed.
Output: C:\\broker-crm\\nda_missing_from_backoffice.xlsx (June + July sheets)."""
import datetime
import openpyxl
from openpyxl.styles import Font
import db_config

conn = db_config.connect(); cur = conn.cursor()

def excel_persons(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    logins = [int(r[2]) for r in list(ws.iter_rows(values_only=True))[1:]
              if r and isinstance(r[2], (int, float))]
    cur.execute('SELECT customer_no FROM clients WHERE login = ANY(%s) AND customer_no IS NOT NULL', (logins,))
    return set(r[0] for r in cur.fetchall())

def crm_new(f, tnext):
    cur.execute("""SELECT cu.customer_no, cu.name FROM customers cu WHERE cu.kind='client'
      AND cu.customer_no IN (SELECT customer_no FROM clients WHERE customer_no IS NOT NULL
          AND COALESCE(first_deposit_at,'')<>'' GROUP BY customer_no
          HAVING MIN(first_deposit_at) >= %s AND MIN(first_deposit_at) < %s)
      AND cu.customer_no NOT IN (SELECT customer_no FROM clients
          WHERE COALESCE(user_archived,FALSE) AND customer_no IS NOT NULL)""", (f, tnext))
    return {r[0]: r[1] for r in cur.fetchall()}

def details(cn, f, tnext):
    cur.execute("""SELECT cl.login, min(left(t.tx_date,16)), min(t.method), round(sum(t.amount)::numeric),
                          min(cl.country), min(u.full_name)
      FROM clients cl
      JOIN transactions t ON t.login=cl.login AND t.tx_type='deposit' AND t.tx_date>=%s AND t.tx_date<%s
      LEFT JOIN users u ON u.id = cl.assigned_agent_id
      WHERE cl.customer_no=%s GROUP BY cl.login ORDER BY 2 LIMIT 1""", (f, tnext, cn))
    return cur.fetchone()

out = openpyxl.Workbook(); out.remove(out.active)
HEAD = ['Customer ID', 'Name', 'Trading account', 'First deposit (date time)', 'Method',
        'Deposited in month ($)', 'Country', 'Sales agent']
for label, path, f, tnext in (
        ('June 2026', r'C:\broker-crm\nda_account June 2026.xlsx', '2026-06-01', '2026-07-01'),
        ('July 2026', r'C:\broker-crm\nda_account.xlsx',           '2026-07-01', '2026-07-11')):
    sheet_set = excel_persons(path)
    ours = crm_new(f, tnext)
    missing = [cn for cn in ours if cn not in sheet_set]
    ws = out.create_sheet(f'{label} — missed ({len(missing)})')
    ws.append(HEAD)
    for c_ in ws[1]: c_.font = Font(bold=True)
    for cn in sorted(missing):
        d = details(cn, f, tnext)
        if d:
            ws.append([cn, ours[cn], d[0], d[1], d[2], float(d[3]), d[4] or '', d[5] or ''])
        else:
            ws.append([cn, ours[cn], '', '(first_deposit_at set, tx outside window)', '', '', '', ''])
    for col, w in zip('ABCDEFGH', (12, 30, 14, 20, 16, 18, 18, 20)):
        ws.column_dimensions[col].width = w
    print(f'{label}: {len(missing)} missed by the sheet')

out.save(r'C:\broker-crm\nda_missing_from_backoffice.xlsx')
print('saved C:\\broker-crm\\nda_missing_from_backoffice.xlsx')
conn.close()
