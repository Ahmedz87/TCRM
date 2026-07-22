import openpyxl, sys, re
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
FILES = {
    "2023": "C:/Broker-crm/IB setting/Commission Report 2023.xlsx",
    "2024": "C:/Broker-crm/IB setting/Commission Report 2024.xlsx",
    "2025": "C:/Broker-crm/IB setting/Commission Report 2025.xlsx",
    "2026": "C:/Broker-crm/IB setting/Commission Report 26.xlsx",
}
TARGET = "1001"   # Mahmood ext_ib_id
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0
def wal(s):
    m = re.match(r"\s*(\d+)", str(s or "")); return m.group(1) if m else None
grand_ag = grand_det = 0.0
for yr, path in FILES.items():
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # Aggregated TotalComm for TARGET (may appear in >1 row)
    agname = next(s for s in wb.sheetnames if s.lower() == "aggregated")
    ag = wb[agname]; hdr = [c.value for c in next(ag.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(hdr)}
    ag_tot = 0.0; ag_rows = 0
    for r in ag.iter_rows(min_row=2, values_only=True):
        if wal(r[idx["Wallet"]]) == TARGET:
            ag_tot += num(r[idx["TotalComm"]]); ag_rows += 1
    # Details(+Details2) per-trade Commission for TARGET
    det = 0.0; det_rows = 0; sheets = [s for s in wb.sheetnames if s.lower().startswith("details")]
    for sh in sheets:
        for r in wb[sh].iter_rows(min_row=2, values_only=True):
            if wal(r[1]) == TARGET:
                det += num(r[9]); det_rows += 1
    wb.close()
    print(f"{yr}: Aggregated TotalComm=${ag_tot:,.2f} ({ag_rows} row) | Details sum=${det:,.2f} ({det_rows} trades, sheets={sheets})")
    grand_ag += ag_tot; grand_det += det
print(f"\nTOTAL Aggregated=${grand_ag:,.2f}   TOTAL Details=${grand_det:,.2f}")
