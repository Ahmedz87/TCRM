import openpyxl, sys, re
from datetime import datetime
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
def parse_dt(s):
    s = re.sub(r"\s+", " ", str(s or "").strip())
    for fmt in ("%b %d %Y %I:%M%p", "%b %d %Y %I:%M %p"):
        try: return datetime.strptime(s, fmt)
        except ValueError: pass
    return None
def scan(path, which):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lo = hi = None; n = 0
    for sh in wb.sheetnames:
        if not sh.lower().startswith("details"): continue
        for r in wb[sh].iter_rows(min_row=2, values_only=True):
            dt = parse_dt(r[0]); n += 1
            if dt:
                if lo is None or dt < lo: lo = dt
                if hi is None or dt > hi: hi = dt
    wb.close()
    print(f"{which}: rows={n}  MIN={lo}  MAX={hi}")
    return lo, hi
scan("C:/Broker-crm/IB setting/Commission Report 2023.xlsx", "2023")
scan("C:/Broker-crm/IB setting/Commission Report 26.xlsx", "2026")
