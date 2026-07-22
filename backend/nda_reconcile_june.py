"""Reconcile the back office's June NDA sheet vs the CRM June new-clients set."""
import datetime
import openpyxl
import db_config

wb = openpyxl.load_workbook(r'C:\broker-crm\nda_account June 2026.xlsx', read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
entries = [r for r in rows[1:] if r and isinstance(r[2], (int, float))]
logins = [int(r[2]) for r in entries]
ds = [r[8] for r in entries if isinstance(r[8], datetime.datetime)]
print(f"sheet rows={len(rows)} clean entries={len(entries)} distinct logins={len(set(logins))}")
if ds:
    print("date range:", min(ds).date(), "->", max(ds).date())

conn = db_config.connect(); cur = conn.cursor()
cur.execute('SELECT login, customer_no FROM clients WHERE login = ANY(%s)', (list(set(logins)),))
m = {r[0]: r[1] for r in cur.fetchall()}
missing = [lg for lg in set(logins) if lg not in m]
excel_persons = set(v for v in m.values() if v)
print(f"logins found in CRM: {len(m)} | not in CRM: {missing}")
print(f"excel distinct persons: {len(excel_persons)}")

# CRM June new clients (people, earliest first deposit in June, clients-page universe)
cur.execute("""SELECT cu.customer_no, cu.name FROM customers cu WHERE cu.kind='client'
  AND cu.customer_no IN (SELECT customer_no FROM clients WHERE customer_no IS NOT NULL AND COALESCE(first_deposit_at,'')<>''
      GROUP BY customer_no HAVING MIN(first_deposit_at) >= '2026-06-01' AND MIN(first_deposit_at) < '2026-07-01')
  AND cu.customer_no NOT IN (SELECT customer_no FROM clients WHERE COALESCE(user_archived,FALSE) AND customer_no IS NOT NULL)""")
ours = {r[0]: r[1] for r in cur.fetchall()}
print(f"CRM June new clients: {len(ours)}")
print(f"overlap: {len(set(ours) & excel_persons)}")

print("=== In CRM, NOT in sheet ===")
for cn in [c for c in ours if c not in excel_persons]:
    cur.execute("""SELECT cl.login, min(left(t.tx_date,10)), min(t.method), round(sum(t.amount)::numeric)
      FROM clients cl JOIN transactions t ON t.login=cl.login AND t.tx_type='deposit'
        AND t.tx_date>='2026-06-01' AND t.tx_date<'2026-07-01'
      WHERE cl.customer_no=%s GROUP BY cl.login LIMIT 1""", (cn,))
    d = cur.fetchone()
    if d:
        print(f"  {cn} | {ours[cn]} | login {d[0]} | {d[1]} | {d[2]} | ${d[3]}")
    else:
        print(f"  {cn} | {ours[cn]} | first_deposit_at in June but NO June deposit tx")

print("=== In sheet, NOT counted by CRM ===")
for cn in excel_persons - set(ours):
    cur.execute("SELECT kind FROM customers WHERE customer_no=%s", (cn,))
    k = (cur.fetchone() or ['?'])[0]
    cur.execute("SELECT MIN(NULLIF(first_deposit_at,'')) FROM clients WHERE customer_no=%s", (cn,))
    fd = cur.fetchone()[0]
    cur.execute("SELECT bool_or(COALESCE(user_archived,false)) FROM clients WHERE customer_no=%s", (cn,))
    ua = cur.fetchone()[0]
    print(f"  {cn} | kind={k} | CRM first deposit: {str(fd)[:16]} | team-archived={ua}")
conn.close()
