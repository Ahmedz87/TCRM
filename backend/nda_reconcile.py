"""Reconcile the back office's nda_account.xlsx (July new clients) against the CRM KPI set."""
import openpyxl
import db_config

wb = openpyxl.load_workbook(r'C:\broker-crm\nda_account.xlsx', read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
entries = [r for r in list(ws.iter_rows(values_only=True))[1:] if r and isinstance(r[2], (int, float))]
logins = [int(r[2]) for r in entries]

conn = db_config.connect(); cur = conn.cursor()
cur.execute('SELECT login, customer_no FROM clients WHERE login = ANY(%s)', (logins,))
excel_persons = set(r[1] for r in cur.fetchall() if r[1])

cur.execute("""SELECT cu.customer_no, cu.name FROM customers cu WHERE cu.kind='client'
  AND cu.customer_no IN (SELECT customer_no FROM clients WHERE customer_no IS NOT NULL AND COALESCE(first_deposit_at,'')<>''
      GROUP BY customer_no HAVING MIN(first_deposit_at) >= '2026-07-01' AND MIN(first_deposit_at) < '2026-07-11')
  AND cu.customer_no NOT IN (SELECT customer_no FROM clients WHERE COALESCE(user_archived,FALSE) AND customer_no IS NOT NULL)""")
ours = {r[0]: r[1] for r in cur.fetchall()}

overlap = set(ours) & excel_persons
extra = [cn for cn in ours if cn not in excel_persons]
print(f"CRM={len(ours)}  excel_persons={len(excel_persons)}  overlap={len(overlap)}")
print("=== In CRM, MISSING from the back-office sheet ===")
for cn in extra:
    cur.execute("""SELECT cl.login, min(left(t.tx_date,10)), min(t.method), round(sum(t.amount)::numeric)
      FROM clients cl JOIN transactions t ON t.login=cl.login AND t.tx_type='deposit'
        AND t.tx_date>='2026-07-01' WHERE cl.customer_no=%s GROUP BY cl.login LIMIT 1""", (cn,))
    d = cur.fetchone()
    if d:
        print(f"  {cn} | {ours[cn]} | login {d[0]} | first dep {d[1]} | {d[2]} | ${d[3]}")
    else:
        print(f"  {cn} | {ours[cn]} | (no July tx found — first_deposit_at only)")
print("=== In sheet, NOT counted by CRM ===")
for cn in excel_persons - set(ours):
    cur.execute("SELECT MIN(NULLIF(first_deposit_at,'')) FROM clients WHERE customer_no=%s", (cn,))
    print(f"  {cn} | CRM first deposit: {cur.fetchone()[0]}")
conn.close()
