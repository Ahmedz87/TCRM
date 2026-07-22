"""
excel_trade_commission — PER-TRADE commission exactly as Plugit paid it, keyed by (login, ticket),
from the 4 Commission Report Details sheets.

DESK RULE (Jul 9 2026): for the excel-covered era the sheets are the ONLY source of per-trade
commission — no recomputation. ib_trades.main() LEFT JOINs this table and overrides commission
for every trade closed on/before EXCEL_CUTOFF. After the cutoff the engine computes normally.

A ticket can appear several times in Details (partial closes) -> SUM per (login, ticket).
Idempotent full rebuild. Re-run when a new yearly report arrives.
"""
import openpyxl, re, sys
from collections import defaultdict
from datetime import datetime
import db_config
from psycopg2.extras import execute_values

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
FILES = [f"C:/Broker-crm/IB setting/Commission Report {y}.xlsx" for y in ("2023", "2024", "2025", "26")]

def ival(s):
    m = re.match(r"\s*(\d+)", str(s or "")); return int(m.group(1)) if m else None
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0

agg = defaultdict(lambda: [0.0, 0.0, None])   # (login, ticket) -> [commission, lots, wallet]
rows = 0
for path in FILES:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    n = 0
    for sh in wb.sheetnames:
        if not sh.lower().startswith("details"): continue
        for r in wb[sh].iter_rows(min_row=2, values_only=True):
            w = ival(r[1]); acc = ival(r[2]); tk = ival(r[3])
            if not acc or not tk: continue
            a = agg[(acc, tk)]
            a[0] += num(r[9]); a[1] += num(r[7]); a[2] = w
            n += 1
    wb.close(); rows += n
    print(f"  {path.split('/')[-1]}: {n:,} rows", flush=True)

print(f"aggregated {len(agg):,} (login, ticket) trades from {rows:,} rows")
con = db_config.connect(); cur = con.cursor()
cur.execute("DROP TABLE IF EXISTS excel_trade_commission")
cur.execute("""CREATE TABLE excel_trade_commission(
    login BIGINT NOT NULL, deal_id BIGINT NOT NULL,
    commission DOUBLE PRECISION NOT NULL, lots DOUBLE PRECISION, ext_ib_id BIGINT,
    PRIMARY KEY (login, deal_id))""")
execute_values(cur,
    "INSERT INTO excel_trade_commission (login, deal_id, commission, lots, ext_ib_id) VALUES %s",
    [(l, t, round(v[0], 6), round(v[1], 4), v[2]) for (l, t), v in agg.items()], page_size=10000)
con.commit()
cur.execute("SELECT COUNT(*), ROUND(SUM(commission)::numeric,0) FROM excel_trade_commission")
print("table rows / total commission:", cur.fetchone())
con.close()
