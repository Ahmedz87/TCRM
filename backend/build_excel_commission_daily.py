"""
excel_commission_daily — per-(IB, day) commission aggregated from the 4 Plugit Commission Reports.

WHY: the IB-profile period KPIs computed commission from our deals recompute, which UNDERCOUNTS the
excel era (missing/indirect trades) — an IB could show period payoff > period commission (impossible:
the wallet is commission-only). This table gives the TRUE per-day commission for 2023-01..EXCEL_CUTOFF;
get_ib sums it for the selected period and adds live ib_trades commission after the cutoff.
Idempotent: full rebuild. Re-run when a new yearly report arrives (and bump ib_trades.EXCEL_CUTOFF).
"""
import openpyxl, re, sys
from collections import defaultdict
from datetime import datetime
import db_config
from psycopg2.extras import execute_values

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
FILES = [f"C:/Broker-crm/IB setting/Commission Report {y}.xlsx" for y in ("2023", "2024", "2025", "26")]

def wal(s):
    m = re.match(r"\s*(\d+)", str(s or "")); return m.group(1) if m else None
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0
def parse_dt(s):
    s = re.sub(r"\s+", " ", str(s or "").strip())
    for fmt in ("%b %d %Y %I:%M%p", "%b %d %Y %I:%M %p"):
        try: return datetime.strptime(s, fmt)
        except ValueError: pass
    return None

agg = defaultdict(float)          # (wallet, 'YYYY-MM-DD') -> commission
total = 0
for path in FILES:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    n = 0
    for sh in wb.sheetnames:
        if not sh.lower().startswith("details"): continue
        for r in wb[sh].iter_rows(min_row=2, values_only=True):
            w = wal(r[1])
            if not w: continue
            dt = parse_dt(r[0]); c = num(r[9])
            if dt and c:
                agg[(int(w), dt.date().isoformat())] += c
                n += 1
    wb.close(); total += n
    print(f"  {path.split('/')[-1]}: {n:,} commission rows", flush=True)

print(f"aggregated: {len(agg):,} (ib, day) pairs from {total:,} rows")
con = db_config.connect(); cur = con.cursor()
cur.execute("""CREATE TABLE IF NOT EXISTS excel_commission_daily(
    ext_ib_id BIGINT NOT NULL, day DATE NOT NULL, commission DOUBLE PRECISION NOT NULL,
    PRIMARY KEY (ext_ib_id, day))""")
cur.execute("TRUNCATE excel_commission_daily")
execute_values(cur, "INSERT INTO excel_commission_daily (ext_ib_id, day, commission) VALUES %s",
               [(w, d, round(c, 4)) for (w, d), c in agg.items()], page_size=5000)
con.commit()
cur.execute("SELECT COUNT(*), ROUND(SUM(commission)::numeric,0) FROM excel_commission_daily")
print("table rows / total commission:", cur.fetchone())
con.close()
