"""Fix ibs.name to match the Plugit Commission Report (agents' data) ClientName per ext_ib_id.
Many IBs showed a generic 'X - IB Account' placeholder; the report has the real full name."""
import openpyxl, re, sys
import db_config
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

def wal(s):
    m = re.match(r"\s*(\d+)", str(s or "")); return m.group(1) if m else None

wb = openpyxl.load_workbook("C:/Broker-crm/IB setting/Commission Report 26.xlsx", read_only=True, data_only=True)
ag = next(s for s in wb.sheetnames if s.lower() == "aggregated"); ws = wb[ag]
hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]; idx = {h: i for i, h in enumerate(hdr)}
names = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    w = wal(r[idx["Wallet"]]); nm = r[idx["ClientName"]]
    if w and nm and str(nm).strip():
        names[int(w)] = str(nm).strip()
wb.close()
print(f"names in Excel: {len(names)}")

con = db_config.connect(); cur = con.cursor()
cur.execute("SELECT id, ext_ib_id, name FROM ibs WHERE ext_ib_id IS NOT NULL")
updated = 0
for ibid, ext, cur_name in cur.fetchall():
    new = names.get(ext)
    if new and new != (cur_name or ""):
        cur.execute("UPDATE ibs SET name = %s WHERE id = %s", (new, ibid))
        updated += 1
con.commit()
print(f"updated {updated} IB names to match the Excel")
cur.execute("SELECT COUNT(*) FROM ibs WHERE name ILIKE '%- IB Account%'")
print("remaining generic '- IB Account' names:", cur.fetchone()[0])
con.close()
