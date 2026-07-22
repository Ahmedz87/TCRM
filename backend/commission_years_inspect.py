import openpyxl, sys, re
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
FILES = {
    "2023": "C:/Broker-crm/IB setting/Commission Report 2023.xlsx",
    "2024": "C:/Broker-crm/IB setting/Commission Report 2024.xlsx",
    "2025": "C:/Broker-crm/IB setting/Commission Report 2025.xlsx",
    "2026": "C:/Broker-crm/IB setting/Commission Report 26.xlsx",
}
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0
for yr, path in FILES.items():
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print(f"\n=== {yr}  sheets={wb.sheetnames} ===")
    agname = next(s for s in wb.sheetnames if s.lower() == "aggregated")
    ag = wb[agname]
    hdr = [c.value for c in next(ag.iter_rows(min_row=1, max_row=1))]
    print("Aggregated cols:", hdr)
    # find col indexes
    idx = {h: i for i, h in enumerate(hdr)}
    di, ii, ti = idx.get("DirectComm"), idx.get("IndirectComm"), idx.get("TotalComm")
    direct = indirect = total = 0.0; n = 0
    for r in ag.iter_rows(min_row=2, values_only=True):
        if r[idx.get("Wallet", 2)] is None: continue
        direct += num(r[di]); indirect += num(r[ii]); total += num(r[ti]); n += 1
    print(f"  IBs={n}  DirectComm=${direct:,.0f}  IndirectComm=${indirect:,.0f}  TotalComm=${total:,.0f}")
    print(f"  Details dims: {wb['Details'].max_row} rows")
    wb.close()
