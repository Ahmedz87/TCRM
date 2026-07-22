"""Reconstruct one IB's monthly level from the 4 Commission Reports, showing the EVIDENCE
(driving symbol + per-lot rate) each month, and mark promotions AND demotions."""
import openpyxl, sys, re
from collections import defaultdict
from datetime import datetime
import db_config
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

TARGET = sys.argv[1] if len(sys.argv) > 1 else "105046"   # ext_ib_id / Wallet
FILES = [f"C:/Broker-crm/IB setting/Commission Report {y}.xlsx" for y in ("2023", "2024", "2025", "26")]

def wal(s):
    m = re.match(r"\s*(\d+)", str(s or "")); return m.group(1) if m else None
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0
def parse_dt(s):
    s = re.sub(r"\s+", " ", str(s or "").strip())
    for fmt in ("%b %d %Y %I:%M%p", "%b %d %Y %I:%M %p"):
        try: return datetime.strptime(s, fmt)
        except ValueError: pass
    return None

cur = db_config.connect().cursor()
cur.execute("SELECT symbol, ib_level, comm_per_lot FROM commission_rates")
ladder = defaultdict(dict)
for sym, lv, cpl in cur.fetchall():
    ladder[sym][int(lv)] = float(cpl or 0)
def tiered(sym):
    d = ladder.get(sym)
    if not d: return False
    return len({round(d[l], 4) for l in (5,6,7,8,9) if l in d}) >= 3
TIERED = {s for s in ladder if tiered(s)}
def infer(sym, rate):
    d = ladder.get(sym); best, bd = None, 1e9
    for lv in (5,6,7,8,9):
        if lv in d and d[lv] > 0:
            diff = abs(d[lv]-rate)/d[lv]
            if diff < bd: best, bd = lv, diff
    return best if bd <= 0.04 else None

mb = defaultdict(lambda: [0.0, 0.0])   # (month, sym) -> [comm, vol]
for path in FILES:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sh in wb.sheetnames:
        if not sh.lower().startswith("details"): continue
        for r in wb[sh].iter_rows(min_row=2, values_only=True):
            if wal(r[1]) != TARGET: continue
            sym = str(r[5] or "").strip()
            if sym not in TIERED: continue
            dt = parse_dt(r[0]); vol = num(r[7]); comm = num(r[9])
            if dt and vol > 0:
                c = mb[(dt.date().isoformat()[:7], sym)]; c[0] += comm; c[1] += vol
    wb.close()

# monthly level = volume-weighted vote across tiered symbols (>=1 lot/mo), + the driving symbol
months = sorted({k[0] for k in mb})
print(f"\nIB {TARGET} — monthly level from the Commission Reports:\n")
prev = None
for m in months:
    votes = defaultdict(float); ev = []
    for (mm, sym), (c, v) in mb.items():
        if mm != m or v < 1.0: continue
        lv = infer(sym, c/v)
        if lv:
            votes[lv] += v
            ev.append((v, sym, round(c/v, 2), lv))
    if not votes: continue
    lvl = max(votes, key=votes.get)
    ev.sort(reverse=True)
    drv = ev[0]
    tag = ""
    if prev is not None and lvl > prev: tag = f"  <== PROMOTION {prev}->{lvl}"
    elif prev is not None and lvl < prev: tag = f"  <== DEMOTION {prev}->{lvl}"
    print(f"  {m}  level IB-{lvl}   (driver: {drv[1]} {drv[0]:.1f} lots @ ${drv[2]}/lot -> L{drv[3]}){tag}")
    prev = lvl
